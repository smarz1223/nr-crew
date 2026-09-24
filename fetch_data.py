"""
NR Crew Fantasy Football - data pipeline
Runs daily on GitHub Actions. Downloads the two published Google Sheets
workbooks, scores every player, reconciles to Weekly Scores, builds
standings (H2H + median), analytics, and league history.

Outputs:
  data/nrcrew_data.json   everything the site needs
  data/recon_report.md    human-readable reconciliation report

Local test: set env vars STATS_FILE and HISTORY_FILE to .xlsx paths.
"""
import io, json, math, os, re, sys, collections, datetime
import openpyxl

# ----------------------------------------------------------------- CONFIG
DRAFT_FILE = "NR_Crew_2026_Draft_Results.xlsx"   # committed to the repo; the draft never changes
STATS_URL = ("https://docs.google.com/spreadsheets/d/e/2PACX-1vRlXVtPeOs3lyaLyj59sBIw4pBSYHHCzVK9JJZIrrF2uaqYaePptKTP0dapJ1GpJi6Pe_2sj0c3atG8/pub?output=xlsx")
HISTORY_URL = ("https://docs.google.com/spreadsheets/d/e/2PACX-1vTs3OiPgWspTYZV9kxXFMaYVps0UhoEHs5pXZNe6ZsSW8SDGGkmQq8aXK3jQ1-zFs_5ARv0p-bm3UgC/pub?output=xlsx")

SEASON = 2026
REG_SEASON_WEEKS = 14          # median game applies weeks 1-14 only
PLAYOFF_WEEKS = [15, 16, 17]
PLAYOFF_TEAMS = 6
BYE_TEAMS = 2
OUT_DIR = "data"

# History names -> 2026 names (2026 names display everywhere)
HISTORY_NAME_MAP = {"Marzella": "MARZ", "Mikey": "MCFADDEN", "Hampson": "KEV"}
# Draft-sheet owner names -> 2026 names
DRAFT_NAME_MAP = {"Mr. Joseph A Gentile": "GENTILE", "meany": "MEANY", "Fig": "FIG",
                  "McFadden": "MCFADDEN", "Russo": "RUSSO", "Kardian": "KARDIAN",
                  "Marzella": "MARZ", "Hampello": "KEV", "Drexello": "DREXELLO",
                  "Danny Free": "FREEMAN"}
DRAFT_PICK_RE = re.compile(r"^(.*?)\((.+?) - (QB|RB|WR|TE|K|DEF)\)$")

# League scoring (from SCORING MODIFIERS tab)
SCORING = {
    # offense
    "Pass Yds": 1 / 25, "Pass TD": 4, "Pass INT": -1,
    "Rush Yds": 1 / 10, "Rush TD": 6,
    "Receptions": 1, "Rec Yds": 1 / 10, "Rec TD": 6,
    "Return TD": 6, "2-Pt Conv": 2, "Fumbles Lost": -2,
    # kicker
    "FG 0-19": 3, "FG 20-29": 3, "FG 30-39": 3, "FG 40-49": 4, "FG 50+": 5,
    "FG Miss 0-19": -1, "FG Miss 20-29": -1, "FG Miss 30-39": -1, "PAT": 1,
    # defense / special teams
    "Sacks": 1, "Def INT": 2, "Fum Rec": 2, "Def TD": 6, "Safety": 2,
    "Blocked Kick": 4, "Def Return TD": 6,
}
DST_PA_TIERS = [15, 10, 5, 2, 0, -1, -5]   # points-allowed tier values

# Column map: (group row label, header row label) -> stat key
POS_COLS = {("Misc", "GP*"): "GP", ("Misc", "2PT"): "2-Pt Conv",
            ("Passing", "Yds"): "Pass Yds", ("Passing", "TD"): "Pass TD", ("Passing", "Int"): "Pass INT",
            ("Rushing", "Yds"): "Rush Yds", ("Rushing", "TD"): "Rush TD",
            ("Receiving", "Rec"): "Receptions", ("Receiving", "Yds"): "Rec Yds", ("Receiving", "TD"): "Rec TD",
            ("Ret", "TD"): "Return TD", ("Fum", "Lost"): "Fumbles Lost"}
K_COLS = {("Misc", "GP*"): "GP",
          ("Field Goals Made", "0-19"): "FG 0-19", ("Field Goals Made", "20-29"): "FG 20-29",
          ("Field Goals Made", "30-39"): "FG 30-39", ("Field Goals Made", "40-49"): "FG 40-49",
          ("Field Goals Made", "50+"): "FG 50+",
          ("Field Goals Missed", "0-19"): "FG Miss 0-19", ("Field Goals Missed", "20-29"): "FG Miss 20-29",
          ("Field Goals Missed", "30-39"): "FG Miss 30-39", ("PAT", "Made"): "PAT"}
DST_COLS = {("Misc", "GP*"): "GP", ("Misc", "Blk Kick"): "Blocked Kick",
            ("Tackles", "Sack"): "Sacks", ("Tackles", "Safe"): "Safety",
            ("Turnovers", "Int"): "Def INT", ("Turnovers", "Fum Rec"): "Fum Rec",
            ("TD", "TD"): "Def TD", ("Ret", "TD"): "Def Return TD"}

# Stat category groupings for the Stat Categories page
CATEGORY_GROUPS = {
    "Passing": ["Pass Yds", "Pass TD", "Pass INT"],
    "Rushing": ["Rush Yds", "Rush TD"],
    "Receiving": ["Receptions", "Rec Yds", "Rec TD"],
    "Misc Offense": ["Return TD", "2-Pt Conv", "Fumbles Lost"],
    "Kicking": ["FG 0-19", "FG 20-29", "FG 30-39", "FG 40-49", "FG 50+",
                "FG Miss 0-19", "FG Miss 20-29", "FG Miss 30-39", "PAT"],
    "Defense": ["Sacks", "Def INT", "Fum Rec", "Def TD", "Safety", "Blocked Kick",
                "Def Return TD", "DST Pts Allowed"],
}
POSITIONS = ["QB", "RB", "WR", "TE", "K", "DEF"]

# Yahoo name-cell cleanup
NOTE_FLAGS = ["No new player Notes", "No new player Note", "New Player Notes", "New Player Note",
              "Player Notes", "Player Note", "Video Forecast"]
STATUS_TAGS = sorted(["IR-R", "IR", "PUP-P", "PUP-R", "NFI-R", "NFI-A", "SUSP", "COVID-19",
                      "DTD", "NA", "O", "Q", "D"], key=len, reverse=True)
TEAM_POS_RE = re.compile(r"([A-Z][A-Za-z]{1,2}) - (QB|WR|RB|TE|K|DEF)\s*$")


# ----------------------------------------------------------------- LOADING
def load_workbook(env_var, url):
    path = os.environ.get(env_var)
    if path:
        return openpyxl.load_workbook(path, data_only=True)
    import requests
    r = requests.get(url, timeout=60)
    r.raise_for_status()
    return openpyxl.load_workbook(io.BytesIO(r.content), data_only=True)


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def r2(x):
    return round(x + 0.0, 2)


def parse_name(raw):
    first = str(raw).split("\n")[0].strip()
    m = TEAM_POS_RE.search(first)
    nfl, pos = (m.group(1).upper(), m.group(2)) if m else (None, None)
    name = first[:m.start()] if m else first
    for f in NOTE_FLAGS:
        name = name.replace(f, "")
    name = name.strip()
    for tag in STATUS_TAGS:  # status glued to the name, e.g. "A.J. BrownIR"
        if name.endswith(tag) and len(name) > len(tag):
            prev = name[-len(tag) - 1]
            if prev.islower() or prev in ".'":
                name = name[:-len(tag)].strip()
                break
    return name, nfl, pos


def read_log_table(ws, colmap, kind):
    """Returns (player_rows, yahoo_totals_by_owner)."""
    rows = list(ws.iter_rows(values_only=True))
    groups, cur = [], None
    for v in rows[0]:
        if v not in (None, ""):
            cur = str(v).strip()
        groups.append(cur)
    headers = [str(v).strip() if v is not None else None for v in rows[1]]
    idx = {}
    for i, (g, h) in enumerate(zip(groups, headers)):
        if (g, h) in colmap:
            idx[colmap[(g, h)]] = i
    missing = set(colmap.values()) - set(idx)
    if missing:
        raise RuntimeError(f"{ws.title}: columns not found {missing}")
    players, totals = [], {}
    for r in rows[2:]:
        owner, name = r[0], r[2]
        if not owner or not name:
            continue
        stats = {k: num(r[i]) for k, i in idx.items()}
        if str(name).strip() == "Totals":
            totals[owner] = stats
            continue
        pname, nfl, pos = parse_name(name)
        players.append({"owner": str(owner).strip(), "player": pname, "nfl": nfl,
                        "pos": pos, "kind": kind, "stats": stats})
    return players, totals


def score(stats):
    pts = {k: stats.get(k, 0) * v for k, v in SCORING.items() if k in stats}
    return pts


# ----------------------------------------------------------------- WEEKLY / STANDINGS
def load_weekly(ws):
    weeks = collections.defaultdict(dict)
    for r in ws.iter_rows(min_row=3, values_only=True):
        wk, team, pf, pa = r[0], r[1], r[2], r[3]
        if wk is None or not team or pf in (None, "") or pa in (None, ""):
            continue
        weeks[int(wk)][str(team).strip()] = {"pf": float(pf), "pa": float(pa)}
    return dict(sorted(weeks.items()))


def pair_opponents(weeks, flags):
    for wk, teams in weeks.items():
        for t, d in teams.items():
            cands = [o for o, e in teams.items() if o != t
                     and abs(e["pf"] - d["pa"]) < 0.005 and abs(e["pa"] - d["pf"]) < 0.005]
            if len(cands) == 1:
                d["opp"] = cands[0]
            else:
                d["opp"] = None
                flags.append(f"Week {wk}: could not uniquely match opponent for {t} ({len(cands)} candidates)")


def result(a, b):
    return "W" if a > b + 1e-9 else ("L" if a < b - 1e-9 else "T")


def norm_cdf(z):
    return 0.5 * (1 + math.erf(z / math.sqrt(2)))


def build_standings(weeks, owners):
    rec = {o: collections.Counter() for o in owners}
    weekly_rows = []
    perf = collections.defaultdict(list)
    for wk, teams in weeks.items():
        scores = sorted((d["pf"] for d in teams.values()), reverse=True)
        n = len(scores)
        median = (scores[n // 2 - 1] + scores[n // 2]) / 2 if n % 2 == 0 else scores[n // 2]
        mean = sum(scores) / n
        sd = math.sqrt(sum((s - mean) ** 2 for s in scores) / n) or 1
        reg = wk <= REG_SEASON_WEEKS
        for t, d in teams.items():
            h2h = result(d["pf"], d["pa"])
            med = result(d["pf"], median) if reg else None
            rank = 1 + sum(1 for s in scores if s > d["pf"] + 1e-9)
            ap = collections.Counter(result(d["pf"], e["pf"]) for o, e in teams.items() if o != t)
            weekly_rows.append({"week": wk, "team": t, "opp": d.get("opp"), "pf": r2(d["pf"]),
                                "pa": r2(d["pa"]), "h2h": h2h, "median": r2(median) if reg else None,
                                "vs_median": med, "rank": rank,
                                "allplay": f"{ap['W']}-{ap['L']}-{ap['T']}",
                                "playoff": not reg})
            if not reg:
                continue
            c = rec[t]
            c["H2H_" + h2h] += 1
            c["MED_" + med] += 1
            c["AP_W"] += ap["W"]; c["AP_L"] += ap["L"]; c["AP_T"] += ap["T"]
            c["PF"] += d["pf"]; c["PA"] += d["pa"]; c["G"] += 1
            perf[t].append(norm_cdf((d["pf"] - mean) / sd) * 100)
    table = []
    for t in owners:
        c = rec[t]
        W = c["H2H_W"] + c["MED_W"]; L = c["H2H_L"] + c["MED_L"]; T = c["H2H_T"] + c["MED_T"]
        games = W + L + T
        h2h_g = c["H2H_W"] + c["H2H_L"] + c["H2H_T"]
        ap_g = c["AP_W"] + c["AP_L"] + c["AP_T"]
        h2h_pct = (c["H2H_W"] + 0.5 * c["H2H_T"]) / h2h_g if h2h_g else 0
        ap_pct = (c["AP_W"] + 0.5 * c["AP_T"]) / ap_g if ap_g else 0
        table.append({
            "team": t, "W": W, "L": L, "T": T,
            "win_pct": round((W + 0.5 * T) / games, 4) if games else 0,
            "h2h": f"{c['H2H_W']}-{c['H2H_L']}-{c['H2H_T']}",
            "median": f"{c['MED_W']}-{c['MED_L']}-{c['MED_T']}",
            "allplay": f"{c['AP_W']}-{c['AP_L']}-{c['AP_T']}",
            "allplay_pct": round(ap_pct, 4),
            "exp_h2h_wins": round(ap_pct * h2h_g, 2),          # all-play expected H2H wins
            "luck": round((h2h_pct - ap_pct + 1) * 50, 1),       # 0-100, 50 = neutral
            "perf_rating": round(sum(perf[t]) / len(perf[t]), 1) if perf[t] else None,
            "PF": r2(c["PF"]), "PA": r2(c["PA"]),
            "PPG": r2(c["PF"] / c["G"]) if c["G"] else 0, "G": c["G"],
        })
    table.sort(key=lambda x: (-x["win_pct"], -x["PF"]))
    for i, row in enumerate(table, 1):
        row["seed"] = i
        row["status"] = "bye" if i <= BYE_TEAMS else ("playoff" if i <= PLAYOFF_TEAMS else "out")
    return table, weekly_rows


# ----------------------------------------------------------------- RECONCILIATION
def achievable_pa_sums(n):
    sums = {0}
    for _ in range(int(n)):
        sums = {s + t for s in sums for t in DST_PA_TIERS}
    return sums


def reconcile(players, yahoo_totals, weeks, owners, flags):
    report = []
    # 1. parsed rows vs Yahoo "Totals" rows
    for kind, tot in yahoo_totals.items():
        for o, ytot in tot.items():
            mine = collections.Counter()
            for p in players:
                if p["owner"] == o and p["kind"] == kind:
                    mine.update(p["stats"])
            for k, v in ytot.items():
                if abs(mine[k] - v) > 0.001:
                    flags.append(f"{o} {kind} table: parsed {k}={mine[k]} vs Yahoo Totals row {v}")
    # 2. freshness: game log vs weeks entered
    weeks_entered = len(weeks)
    max_gp = max((p["stats"].get("GP", 0) for p in players), default=0)
    state = "OK"
    if max_gp > weeks_entered:
        state = "LIVE"
        flags.append(f"Game log shows {int(max_gp)} games but Weekly Scores has {weeks_entered} weeks. "
                     "Week in progress; reconciliation will settle once scores are entered.")
    elif max_gp < weeks_entered:
        state = "STALE"
        flags.append(f"Game log shows only {int(max_gp)} games but Weekly Scores has {weeks_entered} weeks. "
                     "IMPORTHTML may not have refreshed.")
    # 3. per-owner: residual must be a valid DST points-allowed total
    dst_pa = {}
    for o in owners:
        wk_total = sum(w[o]["pf"] for w in weeks.values() if o in w)
        calc = sum(p["points"] for p in players if p["owner"] == o)
        dst_gp = sum(p["stats"].get("GP", 0) for p in players if p["owner"] == o and p["kind"] == "DST")
        resid = wk_total - calc
        ok = abs(resid - round(resid)) < 0.011 and round(resid) in achievable_pa_sums(dst_gp)
        if state == "OK" and not ok:
            flags.append(f"{o}: residual {resid:.2f} is not a valid DST points-allowed total "
                         f"for {int(dst_gp)} DST games")
        dst_pa[o] = round(resid)
        report.append({"team": o, "weekly_total": r2(wk_total), "calculated": r2(calc),
                       "dst_pts_allowed": round(resid), "dst_games": int(dst_gp),
                       "check": "PASS" if ok else ("PENDING" if state != "OK" else "FAIL")})
    overall = "FAIL" if any(r["check"] == "FAIL" for r in report) or \
        any("Totals row" in f or "opponent" in f for f in flags) else state
    if overall == "OK":
        overall = "PASS"
    return report, dst_pa, overall, weeks_entered


# ----------------------------------------------------------------- ANALYTICS
def build_analytics(players, dst_pa, owners):
    raw = {o: collections.Counter() for o in owners}
    pts = {o: collections.Counter() for o in owners}
    by_pos = {o: collections.Counter() for o in owners}
    for p in players:
        o = p["owner"]
        raw[o].update({k: v for k, v in p["stats"].items() if k != "GP"})
        pts[o].update(p["pts_by_cat"])
        by_pos[o][p["pos"]] += p["points"]
    for o in owners:
        pts[o]["DST Pts Allowed"] += dst_pa.get(o, 0)
        by_pos[o]["DEF"] += dst_pa.get(o, 0)
    cats = {}
    for group, keys in CATEGORY_GROUPS.items():
        cats[group] = {o: {"raw": {k: raw[o].get(k, 0) for k in keys if k != "DST Pts Allowed"},
                           "pts": {k: r2(pts[o].get(k, 0)) for k in keys},
                           "total": r2(sum(pts[o].get(k, 0) for k in keys))} for o in owners}
    positions = {o: {pos: r2(by_pos[o].get(pos, 0)) for pos in POSITIONS} for o in owners}
    return cats, positions


# ----------------------------------------------------------------- DRAFT
def build_draft(wb, players, owners, flags):
    """Join draft picks to season scoring. Returns picks + positional value."""
    ws = wb.worksheets[0]
    picks, rnd = [], None
    for r in ws.iter_rows(values_only=True):
        a = r[0]
        if isinstance(a, str) and a.strip().lower().startswith("round"):
            rnd = int(re.sub(r"\D", "", a) or 0)
            continue
        if a in (None, "") or r[1] in (None, ""):
            continue
        m = DRAFT_PICK_RE.match(str(r[1]).strip())
        if not m:
            flags.append(f"Draft: could not read pick '{r[1]}'")
            continue
        raw = str(r[2]).strip()
        owner = DRAFT_NAME_MAP.get(raw, raw)
        if owner not in owners:
            flags.append(f"Draft owner '{raw}' not found in Weekly Scores")
        picks.append({"round": rnd, "pick": int(a), "player": m.group(1).strip(),
                      "nfl": m.group(2).upper(), "pos": m.group(3), "owner": owner})
    if not picks:
        return None
    teams_per_round = max(p["pick"] for p in picks)
    for p in picks:
        p["overall"] = (p["round"] - 1) * teams_per_round + p["pick"]

    # season points by player (starter points; a drafted player with no data scored none for anyone)
    idx = {}
    for pl in players:
        idx.setdefault((pl["player"].lower(), pl["pos"]), []).append(pl)
        idx.setdefault((pl["player"].lower(), None), []).append(pl)
    for p in picks:
        hit = idx.get((p["player"].lower(), p["pos"])) or idx.get((p["player"].lower(), None)) or []
        p["points"] = r2(sum(h["points"] for h in hit))
        p["gp"] = max((h["stats"].get("GP", 0) for h in hit), default=0)
        p["rostered_by"] = sorted({h["owner"] for h in hit})
        p["kept"] = p["owner"] in p["rostered_by"]

    # ranks: overall and within position, by points vs. by draft order
    def rank(rows, key, reverse=True):
        """Rank rows. Ranking on points, players with no starts all tie for last,
        so a late pick who never played never looks like a steal."""
        out = {}
        if key == "points":
            scored = [r for r in rows if r["gp"]]
            prev, rk = None, 0
            for i, row in enumerate(sorted(scored, key=lambda x: (-x["points"], x["overall"])), 1):
                if prev is None or abs(row["points"] - prev) > 1e-9:
                    rk, prev = i, row["points"]
                out[row["overall"]] = rk
            for row in rows:
                out.setdefault(row["overall"], len(rows))
            return out
        for i, row in enumerate(sorted(rows, key=lambda x: (-x[key] if reverse else x[key])), 1):
            out[row["overall"]] = i
        return out
    r_pts = rank(picks, "points")
    r_drf = rank(picks, "overall", reverse=False)
    for p in picks:
        p["points_rank"] = r_pts[p["overall"]]
        p["value"] = r_drf[p["overall"]] - r_pts[p["overall"]]   # + = outperformed draft slot
    by_pos = {}
    for pos in POSITIONS:
        rows = [p for p in picks if p["pos"] == pos]
        rp, rd = rank(rows, "points"), rank(rows, "overall", reverse=False)
        for p in rows:
            p["pos_points_rank"] = rp[p["overall"]]
            p["pos_draft_rank"] = rd[p["overall"]]
            p["pos_value"] = p["pos_draft_rank"] - p["pos_points_rank"]
        by_pos[pos] = len(rows)

    # per-team summary
    summary = []
    for o in owners:
        mine = [p for p in picks if p["owner"] == o]
        drafted_pts = sum(p["points"] for p in mine if p["kept"])
        team_pts = sum(pl["points"] for pl in players if pl["owner"] == o)
        summary.append({"team": o, "picks": len(mine),
                        "draft_points": r2(sum(p["points"] for p in mine)),
                        "points_from_own_picks": r2(drafted_pts),
                        "points_from_adds": r2(team_pts - drafted_pts),
                        "hits": sum(1 for p in mine if p["value"] > 0),
                        "misses": sum(1 for p in mine if p["value"] < 0),
                        "still_rostered": sum(1 for p in mine if p["kept"]),
                        "avg_value": round(sum(p["value"] for p in mine) / len(mine), 1) if mine else 0})
    summary.sort(key=lambda x: -x["draft_points"])
    return {"rounds": max(p["round"] for p in picks), "teams_per_round": teams_per_round,
            "picks": picks, "summary": summary, "pos_counts": by_pos}


# ----------------------------------------------------------------- HISTORY
CURRENT_OWNERS = set()


def hname(n):
    """History name -> display name. Current managers use their 2026 name."""
    n = str(n).strip()
    n = HISTORY_NAME_MAP.get(n, n)
    return n.upper() if n.upper() in CURRENT_OWNERS else n


def build_history(wb, flags):
    ws = wb["Data"]
    hdr = [str(v).strip() if v else None for v in next(ws.iter_rows(max_row=1, values_only=True))]
    col = {h: i for i, h in enumerate(hdr) if h}
    seasons = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[col["Years"]] in (None, ""):
            continue
        g = lambda k: num(r[col[k]])
        seasons.append({"year": int(r[col["Years"]]), "manager": hname(r[col["Manager"]]),
                        "active": str(r[col["Status"]]).strip() == "Active",
                        "first": g("1st"), "second": g("2nd"), "third": g("3rd"),
                        "W": g("Wins"), "L": g("Losses"), "T": g("Ties"),
                        "games": g("Games"), "playoffs": g("Playoffs"),
                        "W_median": g("Wins w/ Median") or None, "L_median": g("Loss w/ Median") or None})
    # all-time
    agg = collections.defaultdict(collections.Counter)
    active = {}
    for s in seasons:
        a = agg[s["manager"]]
        a["years"] += 1
        for k in ("first", "second", "third", "W", "L", "T", "games", "playoffs"):
            a[k] += s[k]
        active[s["manager"]] = active.get(s["manager"], False) or s["active"]
    alltime = []
    for m, a in agg.items():
        alltime.append({"manager": m, "active": active[m], "years": a["years"], "games": a["games"],
                        "first": a["first"], "second": a["second"], "third": a["third"],
                        "top2": a["first"] + a["second"], "top3": a["first"] + a["second"] + a["third"],
                        "W": a["W"], "L": a["L"], "T": a["T"],
                        "win_pct": round((a["W"] + 0.5 * a["T"]) / a["games"], 4) if a["games"] else 0,
                        "playoffs": a["playoffs"],
                        "playoff_pct": round(a["playoffs"] / a["years"], 4)})
    alltime.sort(key=lambda x: (not x["active"], -x["win_pct"]))
    # active rankings (1 = best; losses: fewest = best)
    act = [x for x in alltime if x["active"]]
    ranks = {x["manager"]: {} for x in act}
    for k in ("years", "games", "first", "second", "third", "top2", "top3", "W", "L",
              "win_pct", "playoffs", "playoff_pct"):
        for x in act:
            better = sum(1 for y in act if (y[k] < x[k] if k == "L" else y[k] > x[k]))
            ranks[x["manager"]][k] = better + 1
    # yearly grids
    years = sorted({s["year"] for s in seasons})
    wins_grid = {y: {} for y in years}
    playoff_grid = {y: {} for y in years}
    for s in seasons:
        wins_grid[s["year"]][s["manager"]] = s["W"]
        playoff_grid[s["year"]][s["manager"]] = int(s["playoffs"])
    games_by_year = {y: max(s["games"] for s in seasons if s["year"] == y) for y in years}
    median_years = sorted({s["year"] for s in seasons if s["W_median"]})
    # championships (Championships tab rows 3+ = years in order; names come from Data)
    ch = wb["Championships"]
    champ_detail = {}
    rows = [r for r in ch.iter_rows(min_row=3, values_only=True)]
    for y, r in zip(years, rows):
        if r[6] is None:
            break
        champ_detail[y] = {"W": num(r[6]), "L": num(r[7]), "T": num(r[8]),
                           "draft": int(num(r[9])), "size": int(num(r[10]))}
    champs = []
    for y in years:
        c = next((s for s in seasons if s["year"] == y and s["first"] == 1), None)
        ru = next((s for s in seasons if s["year"] == y and s["second"] == 1), None)
        d = champ_detail.get(y, {})
        if c and d and (abs(c["W"] - d["W"]) > 0.01 or abs(c["L"] - d["L"]) > 0.01):
            flags.append(f"History {y}: champion record in Data ({c['W']}-{c['L']}) "
                         f"differs from Championships tab ({d['W']}-{d['L']})")
        champs.append({"year": y, "champion": c["manager"] if c else None,
                       "runner_up": ru["manager"] if ru else None, **d})
    slot_titles = collections.Counter(c.get("draft") for c in champs if c.get("size") == 10 and c.get("draft"))
    # evolution (seat grid + fill colors)
    ev = wb["Evolution"]
    legend = {}
    for c in ev[29]:
        if c.value in ("Original", "New", "Returning", "Departing"):
            legend[(c.fill.fgColor.rgb or "")[-6:]] = c.value
    evolution = []
    for r in ev.iter_rows(min_row=3, max_row=2 + len(years)):
        y = r[1].value
        if not y:
            continue
        seats = []  # keeps the sheet's seat columns, gaps included
        for c in r[2:14]:
            seats.append({"manager": hname(c.value),
                          "type": legend.get((c.fill.fgColor.rgb or "")[-6:], "Continuing")}
                         if c.value else None)
        while seats and seats[-1] is None:
            seats.pop()
        evolution.append({"year": int(y), "seats": seats})
    # awards (recomputed)
    def best(key, rows, rev=True, fmt=lambda v: v):
        if not rows:
            return None
        top = max(r[key] for r in rows) if rev else min(r[key] for r in rows)
        who = [r for r in rows if abs(r[key] - top) < 1e-9]
        return top, who
    awards = []
    def add(title, key, rows, rev=True, label=lambda r: r["manager"], extra=lambda r: "All"):
        res = best(key, rows, rev)
        if res:
            v, who = res
            awards.append({"award": title, "stat": v, "managers": sorted({label(w) for w in who}),
                           "years": sorted({str(extra(w)) for w in who})})
    add("Most Titles", "first", alltime)
    add("Most Wins", "W", alltime)
    add("Most Losses", "L", alltime)
    add("Best Win % (5+ seasons)", "win_pct", [a for a in alltime if a["years"] >= 5])
    add("Worst Win % (5+ seasons)", "win_pct", [a for a in alltime if a["years"] >= 5], rev=False)
    add("Most Playoff Appearances", "playoffs", alltime)
    add("Most Top 2 Finishes", "top2", alltime)
    add("Most Top 3 Finishes", "top3", alltime)
    sp = [dict(s, pct=(s["W"] + 0.5 * s["T"]) / s["games"]) for s in seasons if s["games"]]
    add("Best Single Season (Win %)", "pct", sp, extra=lambda r: r["year"])
    add("Worst Single Season (Win %)", "pct", sp, rev=False, extra=lambda r: r["year"])
    return {"alltime": alltime, "active_ranks": ranks, "years": years,
            "games_by_year": games_by_year, "median_years": median_years,
            "wins_grid": wins_grid, "playoff_grid": playoff_grid,
            "championships": champs, "titles_by_draft_slot": dict(sorted(slot_titles.items())),
            "evolution": evolution, "awards": awards, "seasons": seasons}


# ----------------------------------------------------------------- MAIN
def main():
    flags = []
    swb = load_workbook("STATS_FILE", STATS_URL)
    hwb = load_workbook("HISTORY_FILE", HISTORY_URL)

    weeks = load_weekly(swb["Weekly Scores"])
    owners = list(dict.fromkeys(t for w in weeks.values() for t in w))
    pair_opponents(weeks, flags)

    players, ytot = [], {}
    for sheet, cmap, kind in (("POSITION_TABLE", POS_COLS, "OFF"),
                              ("KICKER_TABLE", K_COLS, "K"), ("DST_TABLE", DST_COLS, "DST")):
        p, t = read_log_table(swb[sheet], cmap, kind)
        players += p
        ytot[kind] = t
    for p in players:
        p["pts_by_cat"] = score(p["stats"])
        p["points"] = sum(p["pts_by_cat"].values())
        if not p["pos"]:
            flags.append(f"Could not read position for '{p['player']}' ({p['owner']})")
    unknown = {p["owner"] for p in players} - set(owners)
    for u in sorted(unknown):
        flags.append(f"Game log owner '{u}' not found in Weekly Scores")

    recon, dst_pa, status, weeks_entered = reconcile(players, ytot, weeks, owners, flags)
    standings, weekly_rows = build_standings(weeks, owners)
    cats, positions = build_analytics(players, dst_pa, owners)
    CURRENT_OWNERS.update(owners)
    history = build_history(hwb, flags)
    draft = None
    draft_path = os.environ.get("DRAFT_FILE", DRAFT_FILE)
    if os.path.exists(draft_path):
        try:
            dwb = openpyxl.load_workbook(draft_path, data_only=True)
            draft = build_draft(dwb, players, owners, flags)
        except Exception as e:
            flags.append(f"Draft workbook could not be read: {e}")
    else:
        flags.append(f"Draft file '{draft_path}' not found in the repo; the Draft page will be hidden")

    player_out = sorted(({"owner": p["owner"], "player": p["player"], "nfl": p["nfl"], "pos": p["pos"],
                          "gp": int(p["stats"].get("GP", 0)), "points": r2(p["points"]),
                          "stats": {k: v for k, v in p["stats"].items() if k != "GP"},
                          "pts_by_cat": {k: r2(v) for k, v in p["pts_by_cat"].items() if v}}
                         for p in players), key=lambda x: -x["points"])

    now = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    out = {"season": SEASON, "updated": now, "weeks_entered": weeks_entered,
           "reg_season_weeks": REG_SEASON_WEEKS, "playoff_weeks": PLAYOFF_WEEKS,
           "owners": owners, "recon_status": status, "recon": recon, "flags": flags,
           "standings": standings, "weekly": weekly_rows, "categories": cats,
           "positions": positions, "players": player_out, "history": history, "draft": draft,
           "scoring": SCORING}
    os.makedirs(OUT_DIR, exist_ok=True)
    with open(os.path.join(OUT_DIR, "nrcrew_data.json"), "w") as f:
        json.dump(out, f, indent=1, default=float)

    lines = [f"# NR Crew Reconciliation: {status}", f"Updated {now} | Weeks entered: {weeks_entered}", "",
             "| Team | Weekly Total | Calculated | DST Pts Allowed | DST Games | Check |",
             "|---|---|---|---|---|---|"]
    for r in recon:
        lines.append(f"| {r['team']} | {r['weekly_total']:.2f} | {r['calculated']:.2f} | "
                     f"{r['dst_pts_allowed']} | {r['dst_games']} | {r['check']} |")
    lines += ["", "## Flags"] + ([f"- {f}" for f in flags] or ["- None"])
    with open(os.path.join(OUT_DIR, "recon_report.md"), "w") as f:
        f.write("\n".join(lines) + "\n")
    print("\n".join(lines))
    # never fail the run on data flags; the site shows the status badge instead


if __name__ == "__main__":
    main()
